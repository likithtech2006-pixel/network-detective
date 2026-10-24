#!/usr/bin/env python3
"""Build the Review 2 / Review 3 Word documents from their markdown sources.

The output follows the look of the Review 1 report (Network_Detective_Report.docx):
navy headings, 11 pt body text, tables with a navy header row and light-grey rows.

Usage:
    python3 scripts/build_review_docx.py docs/reviews/review2.md docs/Review_2_Design_and_Implementation.docx

Markdown subset understood:
    % key: value          cover lines (eyebrow, title, subtitle, meta) at the top
    ## Heading            section heading;   ### Heading  sub-heading
    paragraphs            **bold**, *italic*, `code` inline
    - item / 1. item      bullet and numbered lists (numbering restarts per list)
    | a | b |             pipe tables, first row is the header
    ```formula / ```code  a centred formula line, or a monospace code block
    ![caption](path)      a figure (path relative to the markdown file), numbered automatically
    ---pagebreak---       page break
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = RGBColor(0x1E, 0x27, 0x61)
SLATE = RGBColor(0x36, 0x45, 0x4F)
GREY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BODY_FONT = "Calibri"
MONO_FONT = "Consolas"
CONTENT_WIDTH = Inches(6.5)

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)")


def add_runs(paragraph, text: str, size: float | None = None, color: RGBColor | None = None,
             bold: bool | None = None) -> None:
    """Add text with **bold**, *italic* and `code` spans as separate runs."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = MONO_FONT
            run._element.rPr.rFonts.set(qn("w:eastAsia"), MONO_FONT)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
            if bold is not None:
                run.bold = bold
        if size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color


def shade(cell, fill: str) -> None:
    """Set a table cell's background colour."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_borders(table) -> None:
    """Thin single borders on every edge, as in the Review 1 tables."""
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "BFBFBF")
        borders.append(el)
    tbl_pr.append(borders)


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_page_number_footer(section, label: str) -> None:
    """Footer: '<label> · page N'."""
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{label}  ·  page ")
    run.font.size = Pt(9)
    run.font.color.rgb = GREY
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        r = p.add_run()
        r.font.size = Pt(9)
        r.font.color.rgb = GREY
        if kind:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), kind)
            r._element.append(fld)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = text
            r._element.append(instr)


class Builder:
    def __init__(self, source: Path, footer_label: str) -> None:
        self.source = source
        self.doc = Document()
        self.figure = 0
        self.table_no = 0
        section = self.doc.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Inches(1))
        normal = self.doc.styles["Normal"]
        normal.font.name = BODY_FONT
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.12
        add_page_number_footer(section, footer_label)

    # ---------------------------------------------------------------- blocks

    def cover(self, fields: dict[str, list[str]]) -> None:
        for key, size, color, bold, before in (
            ("eyebrow", 11, NAVY, True, 0),
            ("title", 24, NAVY, True, 6),
            ("subtitle", 13, SLATE, False, 4),
        ):
            for value in fields.get(key, []):
                p = self.doc.add_paragraph()
                p.paragraph_format.space_before = Pt(before)
                add_runs(p, value, size=size, color=color, bold=bold)
        for value in fields.get("meta", []):
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, value, size=11, color=SLATE)
        rule = self.doc.add_paragraph()
        p_pr = rule._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:color", "1E2761"), ("w:space", "4")):
            bottom.set(qn(k), v)
        bdr.append(bottom)
        p_pr.append(bdr)

    def heading(self, text: str, level: int) -> None:
        p = self.doc.add_heading(level=level)
        p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.font.size = Pt(15 if level == 1 else 12.5)
        run.font.bold = True
        run.font.color.rgb = NAVY
        run.font.name = BODY_FONT
        keep_with_next(p)

    def paragraph(self, text: str) -> None:
        add_runs(self.doc.add_paragraph(), text)

    def list_item(self, text: str, marker: str) -> None:
        p = self.doc.add_paragraph()
        fmt = p.paragraph_format
        fmt.left_indent = Inches(0.35)
        fmt.first_line_indent = Inches(-0.25)
        fmt.space_after = Pt(3)
        p.add_run(f"{marker}\t")
        fmt.tab_stops.add_tab_stop(Inches(0.35))
        add_runs(p, text)

    def table(self, rows: list[list[str]], caption: str | None) -> None:
        header, body = rows[0], rows[1:]
        cols = len(header)
        size = 9 if cols >= 5 else 10
        table = self.doc.add_table(rows=1 + len(body), cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_borders(table)
        for c, text in enumerate(header):
            cell = table.rows[0].cells[c]
            shade(cell, "1E2761")
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            add_runs(cell.paragraphs[0], text, size=size, color=WHITE, bold=True)
        for r, row in enumerate(body, start=1):
            for c in range(cols):
                cell = table.rows[r].cells[c]
                shade(cell, "F2F2F2" if r % 2 else "FFFFFF")
                cell.paragraphs[0].paragraph_format.space_after = Pt(0)
                add_runs(cell.paragraphs[0], row[c] if c < len(row) else "", size=size)
        # Column widths in proportion to their longest text (Word otherwise splits evenly)
        longest = [
            max(len(re.sub(r"[*`]", "", (row[c] if c < len(row) else ""))) for row in rows)
            for c in range(cols)
        ]
        weights = [min(max(n, 6), 45) for n in longest]
        total = sum(weights)
        table.autofit = False
        for row in table.rows:
            for c, cell in enumerate(row.cells):
                cell.width = int(CONTENT_WIDTH * weights[c] / total)
        # Repeat the header row on every page the table spans
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        tr_pr.append(repeat)
        if caption:
            self.table_no += 1
            p = self.doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            add_runs(p, f"Table {self.table_no}: {caption}", size=9, color=GREY)
            p.runs[0].italic = True
        else:
            self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def formula(self, text: str) -> None:
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(text)
        run.italic = True
        run.font.name = "Cambria Math"
        run.font.size = Pt(12)

    def code(self, lines: list[str]) -> None:
        for i, line in enumerate(lines):
            p = self.doc.add_paragraph()
            fmt = p.paragraph_format
            fmt.space_after = Pt(0)
            fmt.space_before = Pt(4) if i == 0 else Pt(0)
            fmt.line_spacing = 1.0
            fmt.left_indent = Inches(0.15)
            p_pr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "F2F2F2")
            p_pr.append(shd)
            run = p.add_run(line if line else " ")
            run.font.name = MONO_FONT
            run._element.rPr.rFonts.set(qn("w:eastAsia"), MONO_FONT)
            run.font.size = Pt(8.5)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def figure_block(self, caption: str, path: str) -> None:
        image = (self.source.parent / path).resolve()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        keep_with_next(p)
        p.add_run().add_picture(str(image), width=CONTENT_WIDTH)
        self.figure += 1
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(10)
        add_runs(cap, f"Figure {self.figure}: {caption}", size=9, color=GREY)
        for run in cap.runs:
            run.italic = True

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---------------------------------------------------------------- parsing

    def build(self, text: str) -> None:
        lines = text.splitlines()
        cover: dict[str, list[str]] = {}
        i = 0
        while i < len(lines) and (lines[i].startswith("%") or not lines[i].strip()):
            if lines[i].startswith("%"):
                key, _, value = lines[i][1:].partition(":")
                cover.setdefault(key.strip(), []).append(value.strip())
            i += 1
        if cover:
            self.cover(cover)

        pending_caption: str | None = None
        while i < len(lines):
            line = lines[i]
            stripped = line.strip()
            if not stripped:
                i += 1
                continue
            if stripped == "---pagebreak---":
                self.page_break()
            elif stripped.startswith("Table: "):
                pending_caption = stripped[len("Table: "):]
            elif stripped.startswith("### "):
                self.heading(stripped[4:], 2)
            elif stripped.startswith("## "):
                self.heading(stripped[3:], 1)
            elif stripped.startswith("```"):
                kind = stripped[3:].strip()
                block = []
                i += 1
                while i < len(lines) and not lines[i].strip().startswith("```"):
                    block.append(lines[i])
                    i += 1
                if kind == "formula":
                    for f in block:
                        if f.strip():
                            self.formula(f.strip())
                else:
                    self.code(block)
            elif stripped.startswith("!["):
                m = re.match(r"!\[(.*)\]\((.+)\)", stripped)
                self.figure_block(m.group(1), m.group(2))
            elif stripped.startswith("|"):
                rows = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                        rows.append(cells)
                    i += 1
                self.table(rows, pending_caption)
                pending_caption = None
                continue
            elif re.match(r"^- ", stripped):
                while i < len(lines) and lines[i].strip().startswith("- "):
                    self.list_item(lines[i].strip()[2:], "•")
                    i += 1
                continue
            elif re.match(r"^\d+\. ", stripped):
                n = 0
                while i < len(lines) and re.match(r"^\d+\. ", lines[i].strip()):
                    n += 1
                    self.list_item(re.sub(r"^\d+\. ", "", lines[i].strip()), f"{n}.")
                    i += 1
                continue
            else:
                para = [stripped]
                i += 1
                while i < len(lines) and lines[i].strip() and not re.match(
                    r"^(#|\||-\s|\d+\.\s|!\[|```|Table: |---pagebreak)", lines[i].strip()
                ):
                    para.append(lines[i].strip())
                    i += 1
                self.paragraph(" ".join(para))
                continue
            i += 1


def main() -> None:
    source, target = Path(sys.argv[1]), Path(sys.argv[2])
    label = sys.argv[3] if len(sys.argv) > 3 else "Network Detective"
    builder = Builder(source, label)
    builder.build(source.read_text(encoding="utf-8"))
    builder.doc.save(target)
    print(f"wrote {target} ({builder.figure} figures, {builder.table_no} captioned tables)")


if __name__ == "__main__":
    main()
